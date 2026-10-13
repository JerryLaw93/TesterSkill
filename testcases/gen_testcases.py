# -*- coding: utf-8 -*-
"""生成测试点思维导图(.xmind/.drawio)与测试用例表(.xlsx)"""
import json, zipfile, os

OUT = os.path.dirname(os.path.abspath(__file__))

# ===================== 测试点树 =====================
# (L1, [(L2, [(L3, [testpoints])])])
TREE = [
    ("登录模块", [
        ("表单校验", [
            ("基础校验", [
                "[异常] 账号为空时点击登录，提示“请输入账号”",
                "[异常] 密码为空时点击登录，提示“请输入密码”",
                "[异常] 验证码为空时点击登录，提示“请输入验证码”",
            ]),
            ("输入边界", [
                "[边界] 账号输入框 maxlength=20，输入超过20位被截断",
                "[边界] 密码输入框 maxlength=32，输入超过32位被截断",
                "[边界] 验证码输入框 maxlength=4，仅能输入4位",
            ]),
        ]),
        ("验证码功能", [
            ("校验逻辑", [
                "[正向] 输入正确验证码（正确账号密码）登录成功",
                "[异常] 输入错误验证码，提示“验证码错误”并自动刷新验证码、清空输入",
                "[正向] 验证码大小写不敏感（小写输入可匹配大写验证码）",
                "[边界] 验证码输入含空格，trim 后参与比对",
            ]),
            ("刷新交互", [
                "[正向] 点击验证码图片刷新出新验证码",
                "[UI] 验证码图片带噪点/干扰线，提示“看不清？点击图片刷新”",
            ]),
        ]),
        ("登录流程", [
            ("正向登录", [
                "[正向] 使用默认账号 admin/123456 登录成功进入用户管理页",
                "[正向] 登录成功后欢迎语显示“你好，{姓名}”，登录表单密码/验证码被清空",
                "[UI] 登录成功无 error 提示，页面切换到 appView",
            ]),
            ("异常登录", [
                "[异常] 账号不存在，提示“账号或密码错误”并刷新验证码",
                "[异常] 密码错误，提示“账号或密码错误”并刷新验证码",
                "[异常] 账号与密码均错误，同样提示“账号或密码错误”",
            ]),
            ("权限校验", [
                "[权限] 已禁用账号（status=0）登录被拦截，提示“该账号已被禁用”",
                "[安全] 登录提示不区分“账号不存在”与“密码错误”，防止账号枚举",
            ]),
        ]),
        ("页面展示", [
            ("UI检查", [
                "[UI] 登录页标题/副标题/默认账号提示（admin/123456）正常显示",
                "[UI] 提示消息红/绿色样式区分 error 与 success",
                "[兼容性] 不同浏览器（Chrome/Safari/Firefox）登录页布局正常",
            ]),
        ]),
    ]),
    ("用户管理模块", [
        ("用户列表", [
            ("列表渲染", [
                "[正向] 登录后用户表格正常渲染（账号/姓名/角色/状态/创建时间/操作）",
                "[正向] 角色标签样式区分：管理员蓝色、普通用户绿色",
                "[正向] 状态点区分启用（绿）/禁用（灰）",
                "[边界] 用户数据为空时表格显示“暂无数据”",
                "[边界] 姓名为空时显示“-”占位",
            ]),
            ("安全显示", [
                "[安全] 账号/姓名含 <script> 等特殊字符时被转义，不执行XSS",
            ]),
        ]),
        ("搜索与筛选", [
            ("关键字搜索", [
                "[正向] 输入账号关键字，实时过滤列表",
                "[正向] 输入姓名关键字，实时过滤列表",
                "[边界] 关键字大小写不区分（大小写混合可匹配）",
                "[异常] 输入无匹配关键字，列表显示“暂无数据”",
                "[边界] 关键字首尾带空格，trim 后仍能匹配",
            ]),
            ("角色筛选", [
                "[正向] 角色筛选“全部角色”显示所有用户",
                "[正向] 筛选“管理员”仅显示 admin 角色用户",
                "[正向] 筛选“普通用户”仅显示 user 角色用户",
                "[正向] 搜索关键字 + 角色筛选组合条件过滤",
            ]),
        ]),
        ("新增用户", [
            ("表单校验", [
                "[异常] 账号含特殊字符或中文，提示“账号需为3-20位字母数字下划线”",
                "[边界] 账号合法边界：3位通过、2位被拒、20位通过、21位被拒",
                "[异常] 密码长度<6位或>32位，提示“密码长度需为6-32位”",
                "[边界] 密码合法边界：6位通过、32位通过、5位/33位被拒",
                "[异常] 已存在的账号重复新增，提示“该账号已存在”",
            ]),
            ("新增流程", [
                "[正向] 填写合法信息新增普通用户成功，列表即时刷新",
                "[正向] 新增管理员角色用户成功",
                "[正向] 新增时状态选择“禁用”，保存后列表显示禁用",
                "[正向] 新增成功后使用新账号可正常登录",
                "[边界] 姓名留空允许保存，列表显示“-”",
            ]),
        ]),
        ("编辑用户", [
            ("编辑流程", [
                "[正向] 编辑姓名/角色/状态保存成功，列表数据更新",
                "[正向] 编辑弹窗回显用户原有数据",
                "[正向] 点击“取消”或遮罩层关闭弹窗，数据不变",
            ]),
            ("字段保护", [
                "[权限] 编辑时账号输入框被禁用不可修改",
                "[权限] 编辑内置 admin 选择禁用状态，提示“内置管理员 admin 不能被禁用”",
                "[正向] 修改用户密码后，旧密码登录失败、新密码登录成功",
                "[正向] 编辑普通用户角色为管理员后角色标签更新",
            ]),
        ]),
        ("删除用户", [
            ("删除流程", [
                "[正向] 点击删除弹出确认框，确认后用户被移除且列表刷新",
                "[正向] 删除确认框点击取消，用户保留",
                "[异常] 删除的用户账号再登录，提示“账号或密码错误”",
            ]),
            ("删除保护", [
                "[权限] 删除内置 admin 时提示“内置管理员 admin 不可删除”",
            ]),
        ]),
    ]),
    ("会话与数据", [
        ("会话管理", [
            ("退出登录", [
                "[正向] 点击退出登录返回登录页，提示“已安全退出”，验证码刷新",
                "[异常] 退出后刷新浏览器不再自动进入用户管理页",
            ]),
            ("会话恢复", [
                "[正向] 登录后刷新页面，通过 sessionStorage 自动恢复登录态",
                "[异常] 登录后账号被禁用，刷新页面不再恢复进入管理页",
                "[边界] sessionStorage 中会话数据非法（被篡改），页面停留登录页不报错",
            ]),
        ]),
        ("数据持久化", [
            ("localStorage", [
                "[正向] 新增/编辑/删除用户后刷新页面数据仍保留",
                "[边界] localStorage 数据被清空后重开页面，恢复默认 admin 账号",
                "[安全] localStorage 数据被篡改为非法 JSON，页面不崩溃并重建默认数据",
                "[性能] 用户量较大（如100条）时列表渲染与搜索无明显卡顿",
            ]),
        ]),
    ]),
]

ROOT_TITLE = "用户管理系统 测试点思维导图"

# ===================== 1. 生成 XMind =====================
def _id():
    i = 0
    while True:
        i += 1
        yield f"t{i}"
idg = _id()

def to_topic(title, kids):
    """kids: list of topic-dict（已构建）或 (title, subkids) 元组"""
    t = {"id": next(idg), "title": title}
    attached = []
    for k in kids:
        if isinstance(k, dict):
            attached.append(k)
        else:
            sub_title, sub_kids = k
            attached.append(to_topic(sub_title, sub_kids or []))
    if attached:
        t["children"] = {"attached": attached}
    return t

tree_topics = []
for l1, l2list in TREE:
    l2_topics = []
    for l2, l3list in l2list:
        l3_topics = []
        for l3, tps in l3list:
            l3_topics.append(to_topic(l3, [(tp, None) for tp in tps]))
        l2_topics.append(to_topic(l2, l3_topics))
    tree_topics.append(to_topic(l1, l2_topics))

content = [{
    "id": "sheet-1", "class": "sheet", "title": "Sheet 1",
    "rootTopic": {"id": "root", "title": ROOT_TITLE, "children": {"attached": tree_topics}}
}]
metadata = {"creator": {"name": "Project Test Team", "version": "1.0.0"}, "version": "R3.7.9"}
manifest = {"file-entries": {"content.json": {}, "metadata.json": {}}}

xmind_path = os.path.join(OUT, "测试点思维导图.xmind")
with zipfile.ZipFile(xmind_path, "w", zipfile.ZIP_DEFLATED) as z:
    z.writestr("content.json", json.dumps(content, ensure_ascii=False))
    z.writestr("metadata.json", json.dumps(metadata, ensure_ascii=False))
    z.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False))
with zipfile.ZipFile(xmind_path) as z:
    assert set(z.namelist()) >= {"content.json", "metadata.json", "manifest.json"}, z.namelist()
print("xmind OK:", xmind_path)

# ===================== 2. 生成 Draw.io =====================
COLORS = {"正向": "#9ACD32", "边界": "#FFB347", "异常": "#FF6B6B", "权限": "#6FA8DC", "性能": "#7FDBFF", "安全": "#B39DDB", "兼容性": "#A2D9A2", "UI": "#F7C8E0"}

mx_cells, edges = [], []
node_id = [1]
W, H = 220, 34
GAPX, GAPY = 60, 8

def new_node(label, x, y, color=None):
    nid = f"n{node_id[0]}"; node_id[0] += 1
    style = "rounded=1;whiteSpace=wrap;html=1;fontSize=11;"
    if color:
        style += f"fillColor={color};strokeColor=#666666;"
    mx_cells.append({"id": nid, "label": label, "x": x, "y": y, "style": style})
    return nid

cursor_y = [40]
def layout(title, kids, level, parent_id):
    x = 40 + level * (W + GAPX)
    if not kids:
        y = cursor_y[0]; cursor_y[0] += H + GAPY
        label = title
        color = None
        for k, c in COLORS.items():
            if title.startswith(f"[{k}]"):
                color = c
                label = title
                break
        nid = new_node(label, x, y, color)
        if parent_id:
            edges.append((parent_id, nid))
        return y
    first_y = cursor_y[0]
    child_ys = [layout(k, sub, level + 1, None) for k, sub in kids]
    # placeholder pass: need parent before children edges; do two-pass
    return None

# two-pass layout: first measure leaves, then create parents
leaf_positions = {}
counter = [0]

def build(title, kids, level):
    x = 40 + level * (W + GAPX)
    if not kids:
        y = 40 + counter[0] * (H + GAPY)
        counter[0] += 1
        color = None
        for k, c in COLORS.items():
            if title.startswith(f"[{k}]"):
                color = c; break
        nid = new_node(title, x, y, color)
        return nid, y
    child_ids, child_ys = [], []
    for item in kids:
        if isinstance(item, str):
            cid, cy = build(item, None, level + 1)
        else:
            k, sub = item
            cid, cy = build(k, sub, level + 1)
        child_ids.append(cid); child_ys.append(cy)
    mid = (min(child_ys) + max(child_ys) + H) / 2 - H / 2
    nid = new_node(title, x, mid, "#D5E8D4" if level == 0 else ("#DAE8FC" if level == 1 else "#FFF2CC"))
    for cid in child_ids:
        edges.append((nid, cid))
    return nid, mid

root_id, _ = build(ROOT_TITLE, [(l1, l2list) for l1, l2list in TREE], 0)

mx = ['<mxfile host="app.diagrams.net">']
mx.append('<diagram id="d1" name="测试点思维导图">')
mx.append('<mxGraphModel dx="1200" dy="800" grid="0" gridSize="10" guides="1" tooltips="1" connect="1" arrows="1" fold="1" page="0" pageScale="1" math="0" shadow="0">')
mx.append('<root>')
mx.append('<mxCell id="0"/><mxCell id="1" parent="0"/>')
for c in mx_cells:
    mx.append(f'<mxCell id="{c["id"]}" value="{c["label"]}" style="{c["style"]}" vertex="1" parent="1"><mxGeometry x="{c["x"]}" y="{c["y"]}" width="{W}" height="{H}" as="geometry"/></mxCell>')
for s, t in edges:
    mx.append(f'<mxCell id="e{node_id[0]}" style="edgeStyle=orthogonalEdgeStyle;rounded=0;strokeColor=#999999;fontSize=10;endArrow=none;" edge="1" parent="1" source="{s}" target="{t}"><mxGeometry relative="1" as="geometry"/></mxCell>')
    node_id[0] += 1
mx.append('</root></mxGraphModel></diagram></mxfile>')
drawio_path = os.path.join(OUT, "测试点思维导图.drawio")
with open(drawio_path, "w", encoding="utf-8") as f:
    f.write("".join(mx))
print("drawio OK:", drawio_path)

# ===================== 3. 生成测试用例表 xlsx =====================
CASES = [
    # (模块编号, 一级, 二级, 三级, 编号, 优先级, 分类, 描述, 前置, 步骤, 预期)
    ("M01","登录模块","表单校验","基础校验","TC-LOG-001","高","功能测试","使用默认账号admin/123456及正确验证码登录成功",
     "1、系统已打开登录页；2、存在默认账号 admin（密码123456）；3、验证码已刷新显示",
     "1、输入账号 admin；2、输入密码 123456；3、输入验证码图片上的4位字符；4、点击“登 录”按钮",
     "1、账号输入成功；2、密码输入成功；3、验证码输入成功；4、登录成功，跳转到用户管理页，欢迎语显示“你好，系统管理员”"),
    ("M01","登录模块","表单校验","基础校验","TC-LOG-002","高","功能测试","账号为空时提交登录被拦截",
     "1、系统已打开登录页",
     "1、账号输入框留空；2、输入密码 123456；3、输入验证码；4、点击“登 录”按钮",
     "1、账号框为空；2、密码输入成功；3、验证码输入成功；4、登录失败，提示“请输入账号”（红色），不进入用户管理页"),
    ("M01","登录模块","表单校验","基础校验","TC-LOG-003","高","功能测试","密码为空时提交登录被拦截",
     "1、系统已打开登录页",
     "1、输入账号 admin；2、密码输入框留空；3、输入验证码；4、点击“登 录”按钮",
     "1、账号输入成功；2、密码框为空；3、验证码输入成功；4、登录失败，提示“请输入密码”（红色）"),
    ("M01","登录模块","表单校验","基础校验","TC-LOG-004","高","功能测试","验证码为空时提交登录被拦截",
     "1、系统已打开登录页",
     "1、输入账号 admin；2、输入密码 123456；3、验证码输入框留空；4、点击“登 录”按钮",
     "1、账号输入成功；2、密码输入成功；3、验证码框为空；4、登录失败，提示“请输入验证码”（红色）"),
    ("M01","登录模块","表单校验","输入边界","TC-LOG-005","中","功能测试","账号输入框20位长度上限校验",
     "1、系统已打开登录页",
     "1、在账号框输入21个字符；2、查看账号框实际内容",
     "1、输入被maxlength截断；2、账号框仅保留前20个字符"),
    ("M01","登录模块","表单校验","输入边界","TC-LOG-006","中","功能测试","密码输入框32位长度上限校验",
     "1、系统已打开登录页",
     "1、在密码框输入33个字符；2、查看密码框实际长度",
     "1、输入被maxlength截断；2、密码框仅保留前32个字符"),
    ("M01","登录模块","表单校验","输入边界","TC-LOG-007","中","功能测试","验证码输入框4位长度上限校验",
     "1、系统已打开登录页",
     "1、在验证码框输入5个字符；2、查看验证码框实际内容",
     "1、输入被maxlength截断；2、验证码框仅保留前4个字符"),
    ("M01","登录模块","验证码功能","校验逻辑","TC-LOG-008","高","功能测试","输入错误验证码登录被拦截并自动刷新",
     "1、系统已打开登录页；2、账号密码正确（admin/123456）",
     "1、输入账号 admin；2、输入密码 123456；3、输入与图片不一致的验证码；4、点击“登 录”按钮",
     "1、账号输入成功；2、密码输入成功；3、验证码输入成功；4、登录失败，提示“验证码错误，请重新输入”，验证码图片刷新，验证码框被清空"),
    ("M01","登录模块","验证码功能","校验逻辑","TC-LOG-009","中","功能测试","验证码大小写不敏感校验",
     "1、系统已打开登录页；2、验证码图片显示4位大写字母/数字",
     "1、输入正确账号密码；2、将验证码以全小写形式输入；3、点击“登 录”按钮",
     "1、账号密码输入成功；2、验证码以小写输入；3、登录成功进入用户管理页"),
    ("M01","登录模块","验证码功能","刷新交互","TC-LOG-010","中","功能测试","点击验证码图片刷新验证码",
     "1、系统已打开登录页；2、当前验证码为A",
     "1、点击验证码图片；2、观察验证码图片内容",
     "1、点击响应正常；2、验证码图片刷新为新的4位字符（大概率与旧码不同）"),
    ("M01","登录模块","登录流程","正向登录","TC-LOG-011","中","功能测试","登录成功后表单敏感字段被清空",
     "1、已使用 admin 登录成功进入用户管理页",
     "1、检查密码框内容；2、检查验证码框内容；3、检查欢迎语显示",
     "1、密码框为空；2、验证码框为空；3、欢迎语显示“你好，系统管理员”"),
    ("M01","登录模块","登录流程","异常登录","TC-LOG-012","高","功能测试","账号不存在时登录失败",
     "1、系统已打开登录页；2、系统中不存在账号 nosuchuser",
     "1、输入账号 nosuchuser；2、输入密码 123456；3、输入正确验证码；4、点击“登 录”按钮",
     "1、账号输入成功；2、密码输入成功；3、验证码输入成功；4、登录失败，提示“账号或密码错误”，验证码刷新"),
    ("M01","登录模块","登录流程","异常登录","TC-LOG-013","高","功能测试","密码错误时登录失败",
     "1、系统已打开登录页；2、账号 admin 存在",
     "1、输入账号 admin；2、输入错误密码 wrongpwd；3、输入正确验证码；4、点击“登 录”按钮",
     "1、账号输入成功；2、错误密码输入成功；3、验证码输入成功；4、登录失败，提示“账号或密码错误”，验证码刷新"),
    ("M01","登录模块","登录流程","权限校验","TC-LOG-014","高","功能测试","禁用账号登录被拦截",
     "1、系统中存在状态为禁用的账号 disabled01",
     "1、输入账号 disabled01；2、输入正确密码；3、输入正确验证码；4、点击“登 录”按钮",
     "1、账号输入成功；2、密码输入成功；3、验证码输入成功；4、登录失败，提示“该账号已被禁用，请联系管理员”"),
    ("M01","登录模块","页面展示","UI检查","TC-LOG-015","低","功能测试","登录页元素完整显示",
     "1、系统已打开登录页",
     "1、检查标题与副标题；2、检查账号/密码/验证码三个输入框及占位符；3、检查默认账号提示“默认账号：admin 密码：123456”；4、检查验证码图片带噪点干扰线及提示文字",
     "1、标题显示“系统登录”，副标题显示“请输入账号、密码和验证码登录系统”；2、三个输入框及占位符均正常；3、默认账号提示正常显示；4、验证码图片有噪点/干扰线，下方显示“看不清？点击图片刷新”"),
    ("M02","用户管理模块","用户列表","列表渲染","TC-USER-001","高","功能测试","登录后用户列表正常渲染",
     "1、已使用 admin 登录进入用户管理页；2、系统中存在用户数据",
     "1、查看用户表格；2、检查表头；3、检查每行数据列",
     "1、表格正常渲染所有用户；2、表头包含账号/姓名/角色/状态/创建时间/操作六列；3、每行显示对应数据，admin 行角色为“管理员”、状态为“启用”"),
    ("M02","用户管理模块","用户列表","列表渲染","TC-USER-002","中","功能测试","姓名为空的用户显示占位符",
     "1、已登录进入用户管理页；2、存在姓名为空的普通用户",
     "1、查看该用户行姓名列",
     "1、姓名列显示“-”占位符"),
    ("M02","用户管理模块","用户列表","安全显示","TC-USER-003","中","安全测试","列表对特殊字符做XSS转义",
     "1、已登录进入用户管理页；2、存在账号或姓名含 <script>alert(1)</script> 的用户（可通过localStorage构造）",
     "1、刷新页面查看用户列表；2、观察是否弹出alert或页面异常",
     "1、特殊字符被转义为文本显示，页面正常；2、不执行脚本，无alert弹窗"),
    ("M02","用户管理模块","搜索与筛选","关键字搜索","TC-USER-004","高","功能测试","按账号关键字实时搜索",
     "1、已登录进入用户管理页；2、列表存在多个用户",
     "1、在搜索框输入某用户账号的部分关键字；2、观察列表变化",
     "1、列表实时过滤，仅显示账号包含该关键字的用户；2、每次输入无需点击即可刷新结果"),
    ("M02","用户管理模块","搜索与筛选","关键字搜索","TC-USER-005","中","功能测试","按姓名关键字搜索",
     "1、已登录；2、存在姓名为“张三”的用户",
     "1、在搜索框输入“张三”；2、观察列表",
     "1、输入成功；2、列表仅显示姓名包含“张三”的用户"),
    ("M02","用户管理模块","搜索与筛选","关键字搜索","TC-USER-006","中","功能测试","搜索关键字大小写不敏感",
     "1、已登录；2、存在账号为 TestUser01 的用户",
     "1、在搜索框输入小写 testuser01；2、观察列表",
     "1、输入成功；2、列表能匹配显示账号 TestUser01 的用户"),
    ("M02","用户管理模块","搜索与筛选","关键字搜索","TC-USER-007","中","功能测试","无匹配关键字显示暂无数据",
     "1、已登录进入用户管理页",
     "1、在搜索框输入不存在的关键字 zzzzzz；2、观察列表区域",
     "1、输入成功；2、列表显示“暂无数据”提示"),
    ("M02","用户管理模块","搜索与筛选","角色筛选","TC-USER-008","中","功能测试","角色筛选-全部角色",
     "1、已登录；2、列表同时存在管理员和普通用户",
     "1、点击角色筛选下拉框选择“全部角色”；2、观察列表",
     "1、选择成功；2、列表显示所有角色的用户"),
    ("M02","用户管理模块","搜索与筛选","角色筛选","TC-USER-009","中","功能测试","角色筛选-管理员/普通用户",
     "1、已登录；2、列表同时存在管理员和普通用户",
     "1、选择筛选“管理员”，观察列表；2、选择筛选“普通用户”，观察列表",
     "1、仅显示角色为管理员的用户；2、仅显示角色为普通用户的用户"),
    ("M02","用户管理模块","搜索与筛选","角色筛选","TC-USER-010","中","功能测试","搜索关键字与角色筛选组合过滤",
     "1、已登录；2、存在多个用户",
     "1、输入关键字 a；2、选择角色筛选“普通用户”；3、观察列表",
     "1、关键字输入成功；2、筛选条件选择成功；3、列表仅显示同时满足账号/姓名含a且角色为普通用户的记录"),
    ("M02","用户管理模块","新增用户","表单校验","TC-USER-011","高","功能测试","正常新增普通用户",
     "1、已使用 admin 登录进入用户管理页",
     "1、点击“+ 新增用户”按钮；2、输入账号 testuser01、密码 test123456、姓名 测试用户；3、角色选择“普通用户”，状态选择“启用”；4、点击“保存”按钮",
     "1、弹窗打开，标题为“新增用户”；2、信息输入成功；3、选项选择成功；4、弹窗关闭，列表新增该用户，角色标签为普通用户，状态为启用，创建时间为当前时间"),
    ("M02","用户管理模块","新增用户","表单校验","TC-USER-012","高","功能测试","新增时重复账号被拒绝",
     "1、已登录；2、系统中已存在账号 admin",
     "1、点击“+ 新增用户”；2、输入账号 admin 及合法密码；3、点击“保存”",
     "1、弹窗打开；2、信息输入成功；3、保存被拒绝，提示“该账号已存在”，弹窗不关闭"),
    ("M02","用户管理模块","新增用户","表单校验","TC-USER-013","高","功能测试","账号格式校验（非法字符/长度不足）",
     "1、已登录进入用户管理页",
     "1、点击“+ 新增用户”；2、输入账号 ab（仅2位）及合法密码，点击“保存”；3、修改账号为 a-b@#（含非法字符），点击“保存”",
     "1、弹窗打开；2、保存被拒绝，提示“账号需为 3-20 位字母、数字或下划线”；3、保存被拒绝，同样格式错误提示"),
    ("M02","用户管理模块","新增用户","表单校验","TC-USER-014","中","功能测试","账号合法长度边界校验",
     "1、已登录进入用户管理页",
     "1、新增账号为3位（如 abc）的用户，点击“保存”；2、新增账号为20位的用户，点击“保存”；3、尝试新增账号为21位的用户",
     "1、3位账号保存成功；2、20位账号保存成功；3、21位账号被拒绝（或被maxlength截断为20位）"),
    ("M02","用户管理模块","新增用户","表单校验","TC-USER-015","中","功能测试","密码长度边界校验",
     "1、已登录进入用户管理页",
     "1、新增用户密码为6位，点击“保存”；2、新增用户密码为32位，点击“保存”；3、新增用户密码为5位，点击“保存”；4、新增用户密码为33位",
     "1、6位密码保存成功；2、32位密码保存成功；3、5位密码被拒绝，提示“密码长度需为 6-32 位”；4、33位密码被拒绝或被截断为32位"),
    ("M02","用户管理模块","新增用户","新增流程","TC-USER-016","中","功能测试","新增禁用状态用户",
     "1、已登录进入用户管理页",
     "1、点击“+ 新增用户”；2、输入合法账号密码姓名；3、状态选择“禁用”；4、点击“保存”",
     "1、弹窗打开；2、信息输入成功；3、状态选择成功；4、保存成功，列表中该用户状态显示为“禁用”且状态点为灰色"),
    ("M02","用户管理模块","新增用户","新增流程","TC-USER-017","中","功能测试","新增用户后可正常登录",
     "1、已登录；2、刚新增账号 newuser01（密码 new123456，状态启用）",
     "1、点击“退出登录”；2、输入账号 newuser01、密码 new123456 及正确验证码；3、点击“登 录”",
     "1、返回登录页；2、信息输入成功；3、登录成功进入用户管理页"),
    ("M02","用户管理模块","编辑用户","编辑流程","TC-USER-018","高","功能测试","正常编辑用户信息并保存",
     "1、已登录；2、存在普通用户 testuser01",
     "1、点击 testuser01 行的“编辑”按钮；2、查看弹窗回显数据；3、修改姓名和角色为管理员；4、点击“保存”",
     "1、弹窗打开，标题为“编辑用户”；2、账号/密码/姓名/角色/状态均正确回显；3、修改成功；4、弹窗关闭，列表中该用户姓名更新、角色标签变为管理员"),
    ("M02","用户管理模块","编辑用户","字段保护","TC-USER-019","中","功能测试","编辑时账号字段不可修改",
     "1、已登录；2、列表存在用户",
     "1、点击某用户“编辑”按钮；2、尝试修改账号输入框内容",
     "1、弹窗打开；2、账号输入框为禁用状态（灰色不可编辑）"),
    ("M02","用户管理模块","编辑用户","字段保护","TC-USER-020","高","功能测试","内置admin不能被禁用",
     "1、已登录；2、当前仅存在一个 admin 账号",
     "1、点击 admin 行“编辑”按钮；2、状态选择“禁用”；3、点击“保存”",
     "1、弹窗打开；2、状态选择成功；3、保存被拒绝，提示“内置管理员 admin 不能被禁用”"),
    ("M02","用户管理模块","编辑用户","编辑流程","TC-USER-021","中","功能测试","取消编辑不保存数据",
     "1、已登录；2、存在用户 testuser01",
     "1、点击“编辑”修改姓名为“新姓名”；2、点击“取消”按钮关闭弹窗；3、查看列表",
     "1、修改成功；2、弹窗关闭；3、列表中该用户姓名保持原值未变化"),
    ("M02","用户管理模块","编辑用户","字段保护","TC-USER-022","中","功能测试","修改密码后新旧密码登录验证",
     "1、已登录；2、存在用户 testuser01（密码 test123456）",
     "1、编辑 testuser01 将密码改为 abc12345；2、保存后退出登录；3、使用旧密码 test123456 登录；4、使用新密码 abc12345 登录",
     "1、密码修改成功；2、返回登录页；3、旧密码登录失败，提示“账号或密码错误”；4、新密码登录成功"),
    ("M02","用户管理模块","删除用户","删除流程","TC-USER-023","高","功能测试","确认删除普通用户",
     "1、已登录；2、存在普通用户 testuser01",
     "1、点击 testuser01 行“删除”按钮；2、在确认弹窗中点击“确定”；3、查看列表",
     "1、弹出确认框“确定删除用户「testuser01」吗？”；2、确认成功；3、该用户从列表消失"),
    ("M02","用户管理模块","删除用户","删除流程","TC-USER-024","中","功能测试","取消删除保留用户",
     "1、已登录；2、存在普通用户 testuser01",
     "1、点击“删除”按钮；2、在确认弹窗中点击“取消”；3、查看列表",
     "1、弹出确认框；2、取消成功；3、用户仍在列表中"),
    ("M02","用户管理模块","删除用户","删除保护","TC-USER-025","高","功能测试","内置admin账号不可删除",
     "1、已登录进入用户管理页",
     "1、点击 admin 行的“删除”按钮；2、观察提示",
     "1、不弹出删除确认框；2、弹出提示“内置管理员 admin 不可删除”，admin 仍保留在列表中"),
    ("M02","用户管理模块","删除用户","删除流程","TC-USER-026","中","功能测试","已删除用户无法登录",
     "1、已登录；2、用户 testuser01 已被删除",
     "1、退出登录；2、使用账号 testuser01 及原密码登录；3、观察结果",
     "1、返回登录页；2、信息输入成功；3、登录失败，提示“账号或密码错误”"),
    ("M03","会话与数据","会话管理","退出登录","TC-SES-001","高","功能测试","退出登录返回登录页",
     "1、已使用 admin 登录进入用户管理页",
     "1、点击“退出登录”按钮；2、观察页面切换与提示；3、观察验证码图片",
     "1、返回登录页；2、显示绿色提示“已安全退出”；3、验证码已刷新为新的字符"),
    ("M03","会话与数据","会话管理","会话恢复","TC-SES-002","中","功能测试","刷新页面自动恢复登录态",
     "1、已使用 admin 登录进入用户管理页",
     "1、刷新浏览器页面；2、观察显示的页面",
     "1、页面重新加载；2、无需重新登录，直接进入用户管理页（sessionStorage 会话恢复）"),
    ("M03","会话与数据","会话管理","会话恢复","TC-SES-003","中","功能测试","退出后刷新不恢复登录态",
     "1、已登录后点击退出登录返回登录页",
     "1、刷新浏览器页面；2、观察显示的页面",
     "1、页面重新加载；2、仍停留在登录页，不自动进入用户管理页"),
    ("M03","会话与数据","会话管理","会话恢复","TC-SES-004","中","功能测试","禁用账号刷新后不恢复登录态",
     "1、账号 testuser01 已登录；2、在另一处将该账号状态改为禁用",
     "1、刷新页面；2、观察显示的页面",
     "1、页面重新加载；2、不恢复进入用户管理页，停留在登录页"),
    ("M03","会话与数据","数据持久化","localStorage","TC-SES-005","中","功能测试","用户数据localStorage持久化",
     "1、已登录；2、新增一个用户 persist01",
     "1、刷新浏览器页面；2、查看用户列表",
     "1、页面重新加载并恢复登录态；2、新增的用户 persist01 仍在列表中"),
    ("M03","会话与数据","数据持久化","localStorage","TC-SES-006","低","功能测试","清空localStorage后恢复默认数据",
     "1、已登录；2、开发者工具中清空 localStorage",
     "1、刷新浏览器页面；2、查看用户列表",
     "1、页面重新加载；2、系统重建默认数据，列表显示默认账号 admin"),
    ("M03","会话与数据","数据持久化","localStorage","TC-SES-007","低","安全测试","localStorage数据被篡改不崩溃",
     "1、已打开登录页；2、将 localStorage 中 sys_users 值改为非法JSON字符串",
     "1、刷新浏览器页面；2、观察页面状态与控制台",
     "1、页面正常打开不白屏；2、无JS报错，系统重建默认数据，admin 可正常登录"),
    ("M03","会话与数据","数据持久化","localStorage","TC-SES-008","低","性能测试","较多用户数据下列表渲染与搜索性能",
     "1、已登录；2、通过localStorage构造100条用户数据",
     "1、刷新页面观察列表渲染；2、在搜索框输入关键字实时过滤",
     "1、列表一次性渲染完成无明显卡顿（<1s）；2、搜索过滤响应流畅无卡顿"),
]

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "测试管理平台—测试用例导入模版"

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
GROUP_FILL = PatternFill("solid", fgColor="D9E2F3")
GROUP_FONT = Font(bold=True, size=11)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Border(*[Side(style="thin", color="BFBFBF")]*4)

# Row 2 title
ws.merge_cells("B2:S2")
ws["B2"] = "测试用例"
ws["B2"].font = Font(bold=True, size=16)
ws["B2"].alignment = CENTER

# Row 4 round grouping
for rng, label in [("V4:Y4","第1轮"),("Z4:AC4","第2轮"),("AD4:AG4","第N轮")]:
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.value = label
    c.font = GROUP_FONT
    c.alignment = CENTER
    c.fill = GROUP_FILL

total = len(CASES)
# Rows 5-9 statistics
stats = [
    (5, "用例条数", total, "OK", 0),
    (6, "纳入测试", total, "POK", 0),
    (7, "执行率", "0%", "NG", 0),
    (8, "通过率", "0%", "SKIP", 0),
    (9, "NT", 0, "队列中", total),
]
for r, m1, v1, m2, v2 in stats:
    ws.cell(row=r, column=22, value=m1).font = GROUP_FONT   # V
    ws.cell(row=r, column=23, value=v1)
    ws.cell(row=r, column=24, value=m2).font = GROUP_FONT   # X
    ws.cell(row=r, column=25, value=v2)

# identification block D~H rows 8-9
for col, val in [(4,"测试标识"),(5,"产品"),(6,"系统"),(7,"设计人员"),(8,"可用标签")]:
    ws.cell(row=8, column=col, value=val).font = GROUP_FONT
for col, val in [(4,"TC"),(5,"用户管理系统"),(6,"用户管理系统登录与用户管理"),(7,"AI测试组"),(8,"")]:
    ws.cell(row=9, column=col, value=val)

HEADERS = ["序号","模块编号","一级模块","二级模块","三级模块","用例编号","优先级","分类","是否已评审","用例描述",
           "前置条件","测试方法和步骤","预期结果","已自动化","自动化脚本名称","自动化脚本设计人员姓名","验证过程","测试工具",
           "修订状态","纳入测试","测试结果","测试版本","测试人","备注","测试结果","测试版本","测试人","备注",
           "测试结果","测试版本","测试人","备注"]
# Row 10 headers, starting col B(2)
for i, h in enumerate(HEADERS):
    c = ws.cell(row=10, column=2 + i, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = CENTER
    c.border = THIN

# Data rows
for idx, case in enumerate(CASES, start=1):
    mcode, l1, l2, l3, cid, prio, cat, desc, pre, steps, exp = case
    r = 10 + idx
    row_vals = [idx, mcode, l1, l2, l3, cid, prio, cat, "否", desc, pre, steps, exp, "否", "", "", "", "", "", "是",
                "", "", "", "", "", "", "", "", "", "", ""]
    for i, v in enumerate(row_vals):
        c = ws.cell(row=r, column=2 + i, value=v)
        c.alignment = LEFT_WRAP
        c.border = THIN

# Column widths
widths = {"A":2,"B":6,"C":10,"D":16,"E":16,"F":16,"G":18,"H":8,"I":12,"J":10,"K":40,"L":35,"M":45,"N":45}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
for col in ["O","P","Q","R","S"]:
    ws.column_dimensions[col].width = 16
for i in range(20, 34):  # T..AG
    ws.column_dimensions[get_column_letter(i)].width = 12
ws.column_dimensions["AH"].width = 2

last_row = 10 + total
ws.freeze_panes = "A11"
ws.auto_filter.ref = f"B10:AG{last_row}"

# 统计 sheet
st = wb.create_sheet("统计")
hi = sum(1 for c in CASES if c[5] == "高")
mid = sum(1 for c in CASES if c[5] == "中")
lo = sum(1 for c in CASES if c[5] == "低")
st["A1"] = "测试用例统计"
st["A1"].font = Font(bold=True, size=14)
rows = [
    ("总用例数", total),
    ("优先级-高", hi), ("优先级-中", mid), ("优先级-低", lo),
]
mods = {}
for c in CASES:
    mods[c[1]] = mods.get(c[1], 0) + 1
rows += [(f"模块-{k}", v) for k, v in mods.items()]
cats = {}
for c in CASES:
    cats[c[6]] = cats.get(c[6], 0) + 1
rows += [(f"分类-{k}", v) for k, v in cats.items()]
rows.append(("预估执行时间(分钟)", total * 5))
for i, (k, v) in enumerate(rows, start=3):
    st.cell(row=i, column=1, value=k)
    st.cell(row=i, column=2, value=v)
st.column_dimensions["A"].width = 26
st.column_dimensions["B"].width = 12

xlsx_path = os.path.join(OUT, "测试用例.xlsx")
wb.save(xlsx_path)
print("xlsx OK:", xlsx_path, "cases:", total)
